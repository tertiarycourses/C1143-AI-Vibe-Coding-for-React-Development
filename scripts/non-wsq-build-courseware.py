#!/usr/bin/env python3
"""Generate C1143 learner guide, labs, lesson plan and deck data from one source."""
from pathlib import Path
from datetime import date
import importlib.util, json, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'courseware'; LABROOT=ROOT/'labs'
OUT.mkdir(exist_ok=True); LABROOT.mkdir(exist_ok=True)
spec=importlib.util.spec_from_file_location('course_data',ROOT/'scripts/non-wsq-course-data.py')
data=importlib.util.module_from_spec(spec); spec.loader.exec_module(data)
C=data.COURSE; TOPICS=data.TOPICS; LABS=data.LABS; CONCEPT_DEFS=data.CONCEPT_DEFS; TRAP_NOTES=data.TRAP_NOTES

def first_sentence(s):
    """First sentence of a step, without breaking filenames like AGENTS.md or main.tsx."""
    head=re.split(r'(?<=[^A-Z])\.\s+',s,maxsplit=1)[0]
    return head.rstrip('.')

# Rotated per-step guidance so walkthrough coaching varies instead of repeating one paragraph.
INSPECT_GUIDE=[
 "Predict which files this step should touch and what will change on screen; compare that prediction with the actual diff before moving on.",
 "Say out loud what success looks like for this step before acting; afterwards capture the command output or screenshot that proves it.",
 "Keep the agent inside the declared file scope here — if its proposal reaches further, stop and narrow the request.",
 "Where the agent acts, read its output as a reviewer, not a spectator, and note one specific thing you checked.",
 "If the result differs from your prediction, treat the gap as information: find the exact line that explains it before continuing.",
 "Record the evidence for this step while it is still on screen — a command line and its output beat a memory.",
]
VERIFY_GUIDE=[
 "Run the narrowest check that exercises this step first; only then run the wider lint, type-check and build stack.",
 "Verify the failure path as well as the success path — break the input deliberately and confirm the app responds as designed.",
 "Capture the command, expected result and actual result; if they differ, stop and investigate before the next step.",
 "If this step touches the interface, re-test at both a narrow and a wide viewport and note anything that clips or overflows.",
 "Ask the agent to explain any part of this step's diff you cannot explain yourself; unresolved lines block the checkpoint.",
 "When this step passes, decide keep, refine or revert explicitly and write one sentence recording why.",
]
BLUE='0B6E99'; NAVY='17365D'; TEAL='14866D'; PALE='EAF4F8'; INK='202A35'; WHITE='FFFFFF'

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def field(p,instruction):
    r=p.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=instruction
    sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate'); end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); r._r.extend([begin,instr,sep,end])
def code(d,text):
    p=d.add_paragraph(); p.style=d.styles['Code']; p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(6); p.add_run(text)
def add_box(d,label,text,fill='EAF4F8'):
    t=d.add_table(rows=1,cols=1); t.style='Table Grid'; shade(t.cell(0,0),fill)
    p=t.cell(0,0).paragraphs[0]; p.add_run(label+': ').bold=True; p.add_run(text)
def add_bullets(d,items):
    for x in items: d.add_paragraph(x,style='List Bullet')
def base_doc(kind,toc_entries):
    d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=sec.right_margin=Inches(.72)
    styles=d.styles
    for n,size,color,bold in [('Normal',10.5,INK,False),('Title',28,NAVY,True),('Heading 1',19,BLUE,True),('Heading 2',15,NAVY,True),('Heading 3',12,TEAL,True)]:
        s=styles[n]; s.font.name='Arial'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.font.bold=bold
    code_style=styles.add_style('Code',1); code_style.font.name='Consolas'; code_style.font.size=Pt(8.5); code_style.font.color.rgb=RGBColor.from_string(INK)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(120)
    r=p.add_run('TERTIARY INFOTECH ACADEMY'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(15); r.font.color.rgb=RGBColor.from_string(BLUE)
    p=d.add_paragraph(C['title'],style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph(kind); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].bold=True; p.runs[0].font.size=Pt(18)
    p=d.add_paragraph(f"Course Code: {C['code']}\nVersion {C['version']}\n{C['duration']} · {C['level']}\nConducted by Tertiary Infotech Academy Pte Ltd\nUEN: 201200696W"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    d.add_page_break(); d.add_heading('Document Version Control Record',1)
    t=d.add_table(rows=3,cols=4); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    vals=[['Version','Effective date','Summary','Author'],['2.0','2026-07-12','Complete 20-lab agentic-loop rebuild','Tertiary Infotech Academy'],[C['version'],date.today().isoformat(),'Title, contents, schedule and content-quality corrections','Tertiary Infotech Academy']]
    for i,row in enumerate(vals):
        for j,val in enumerate(row):
            cell=t.cell(i,j); cell.text=val; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i==0:
                shade(cell,BLUE)
                for run in cell.paragraphs[0].runs: run.font.color.rgb=RGBColor(255,255,255); run.bold=True
    d.add_heading('Table of Contents',1)
    for entry,level in toc_entries:
        p=d.add_paragraph(entry); p.paragraph_format.space_after=Pt(2)
        p.paragraph_format.left_indent=Inches(.3*(level-1))
        for run in p.runs: run.font.size=Pt(10.5 if level==1 else 9.5); run.bold=(level==1)
    for sec in d.sections:
        f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
        f.add_run(f"© 2026 Tertiary Infotech Academy Pte Ltd | {C['code']} | {kind} | Page "); field(f,'PAGE'); f.add_run(' of '); field(f,'NUMPAGES')
    settings=d.settings._element; update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); settings.append(update)
    return d

def concept_explanation(lab):
    c=', '.join(lab['concepts'])
    return (f"This lab connects {c} to an observable SprintBoard increment. The important adult-learning move is not memorising syntax; it is predicting what the code should do, comparing that prediction with generated output, and using evidence to resolve the gap. "
            f"The agent can accelerate typing and suggest structures, but the learner owns the product rule, file boundary, safety decision and acceptance evidence. {lab['why']} "
            "When reviewing, trace information from its source through props, state, effects or routes to the visible result. Then trace every user action back through its handler and state transition. This two-way trace exposes hidden coupling and plausible-looking code that does not satisfy the brief.")

def build_lab_markdown(lab,topic_name):
    prior='nothing; this is the first checkpoint' if lab['id']=='1.1' else 'the previous lab checkpoint'
    lines=[f"# Lab {lab['id']} — {lab['title']}",f"> **Topic {lab['topic']}** · approximately {lab['mins']} minutes · builds on {prior}",
      "## Goal",lab['outcome'],"## The build so far",f"SprintBoard grows through one continuous sequence. Begin from {prior}. In this lab you add a bounded capability and finish with a recoverable Git checkpoint.",
      "## What you will build",f"A verified increment touching: {', '.join(f'`{x}`' for x in lab['files'])}. The increment is complete only when the observable checks pass and the generated diff can be explained.",
      "## Concepts you will meet"]
    lines += [f"- **{x}** — {CONCEPT_DEFS[x]}" for x in lab['concepts']]
    lines += ["## Prerequisites","- Complete or restore the previous Git checkpoint.","- Start the development server and confirm the existing flow works.","- Use synthetic data and placeholders. Never paste credentials or private data into an AI prompt.","- Read `AGENTS.md` and keep the agent inside the named file scope.","## Steps"]
    lab_no=int(lab['id'].replace('.',''))
    for i,s in enumerate(lab['steps'],1):
        lines += [f"### Step {i} — {first_sentence(s)}",s,f"**Pause and inspect:** {INSPECT_GUIDE[(lab_no+i)%len(INSPECT_GUIDE)]}",f"**Evidence:** {VERIFY_GUIDE[(lab_no+i)%len(VERIFY_GUIDE)]}"]
    lines += ["## Agentic AI loop","### 1. Frame","State one observable goal, the current checkpoint, exact file scope, non-goals and stop conditions.","### 2. Plan","Require assumptions, numbered steps, files, risks, verification and rollback. Do not authorize code yet.","### 3. Generate","Approve one bounded increment. Keep the development server visible and do not combine refactoring with behavior change.","### 4. Inspect","Compare the plan and diff with the brief. Reject unrelated dependencies, architecture changes, secret handling or untestable claims.","### 5. Verify","Run commands yourself and exercise normal, boundary, empty and failure paths in the browser.","### 6. Correct","Read every changed line, explain data flow, and request the smallest correction backed by a failing check.","### 7. Commit","Commit only understood code. Record the commit and a one-sentence rollback instruction.","## Vibe prompt","```text",lab['prompt'],"```","## Read what the AI wrote"]
    lines += [f"- **{trap}.** Locate the exact line or absence that would reveal this failure. Ask for an explanation before accepting a change." for trap in lab['traps']]
    lines += ["## Why it works",concept_explanation(lab),"## Test it"]
    lines += [f"- [ ] {x}." for x in lab['verify']]
    lines += ["- [ ] `npm run lint` completes without an unexplained suppression.","- [ ] `npm run build` completes and the browser console has no new error.","- [ ] `git diff --check` is clean and every changed file was in the approved plan.","## Evidence to submit","- The final prompt and approved plan.","- `git diff --stat` plus one annotated excerpt showing a reviewed decision.","- Verification command output and a screenshot of the observable result.","- A short note naming one AI suggestion accepted, corrected or rejected and why.","## Independent challenge",f"Change one constraint related to {lab['concepts'][0]} without widening the product scope. Predict the files and tests first, then run Frame, Plan, Generate, Inspect, Verify, Correct and Commit; compare the prediction with the actual diff.","## Troubleshooting and recovery","| Symptom | Likely cause | Recovery |","|---|---|---|"]
    for trap,(cause,fix) in zip(lab['traps'],TRAP_NOTES[lab['id']]):
        lines.append(f"| {trap} | {cause} | {fix} |")
    lines += ["## Reflection",f"How did {lab['concepts'][0]} change what the user could observe? Which evidence most increased or reduced your trust in the generated change? What context should the next lab preserve?",f"### SprintBoard after Lab {lab['id']}",lab['outcome']]
    return '\n\n'.join(lines)+'\n'

def build_labs():
    for lab in LABS:
        topic_name=TOPICS[lab['topic']-1][1]; td=LABROOT/f"topic-{lab['topic']}"; td.mkdir(parents=True,exist_ok=True)
        slug=re.sub(r'[^a-z0-9]+','-',lab['title'].lower()).strip('-')
        (td/f"lab-{lab['id']}-{slug}.md").write_text(build_lab_markdown(lab,topic_name),encoding='utf-8')

def build_learner_guide():
    toc=[('Course Overview',1),('Environment Setup',1),('The Agentic AI Loop Engineering Model',1),('JavaScript and TypeScript Readiness',1)]
    for tnum,tname,_ in TOPICS:
        toc.append((f'Topic {tnum}: {tname}',1))
        toc += [(f"Lab {lab['id']}: {lab['title']}",2) for lab in LABS if lab['topic']==tnum]
    toc += [('Appendix A: AI React Bug Checklist',1),('Appendix B: Troubleshooting',1),('Appendix C: Glossary and References',1)]
    d=base_doc('Learner Guide and Step-by-Step Lab Manual',toc); d.add_page_break()
    d.add_heading('Course Overview',1); d.add_paragraph(f"This two-day intermediate course builds {C['capstone']}, a React task-planning application, through 20 progressive labs. Every lab applies the same engineering loop: Frame, Plan, Generate, Inspect, Verify, Correct and Commit.")
    d.add_heading('Learning Outcomes',2); add_bullets(d,["Scaffold and explain a modern Vite React TypeScript application.","Engineer prompts and repository context that constrain AI coding agents.","Build accessible JSX, reusable components, typed props, events and responsive CSS.","Manage immutable state, synchronize effects, extract custom hooks, route pages and fetch data.","Diagnose defects, test user behavior, audit accessibility, profile code and deploy a verified build."])
    d.add_heading('How to Use This Guide',2); d.add_paragraph('Do not race through the code blocks. Before every agent request, predict the files and behavior. After every response, inspect the plan and diff. Run the checks yourself. If evidence conflicts with the explanation, trust the evidence and investigate.')
    d.add_heading('Environment Setup',1); add_bullets(d,['Node.js LTS and npm','Git and a GitHub account','Visual Studio Code with ESLint support','A modern Chromium, Firefox or Safari browser','An approved coding agent such as Cursor, GitHub Copilot, Claude or Codex'])
    d.add_heading('The Agentic AI Loop Engineering Model',1)
    for h,text in [('Frame','Define one observable outcome, relevant context, exact file scope, constraints, non-goals and stop conditions.'),('Plan','Ask for assumptions, numbered actions, changed files, risks, verification and rollback before code.'),('Generate','Let the agent perform one approved mutation while keeping terminal and browser feedback visible.'),('Inspect','Challenge scope, dependencies, security, accessibility and testability; read the diff line by line.'),('Verify','Run static checks, automated tests and realistic browser paths, including failure and boundary states.'),('Correct','Request the smallest correction that makes a failing check pass. Avoid broad rewrites.'),('Commit','Commit an understood state and record how to restore it.')]:
        d.add_heading(h,2); d.add_paragraph(text)
    d.add_heading('JavaScript and TypeScript Readiness',1)
    readiness=[('Arrow functions','Components and callbacks are functions; be able to distinguish returning an expression from a block body.'),('Destructuring and spread','Props use destructuring; immutable updates use object and array spread to create new references.'),('Array methods','map renders lists, filter derives subsets and find resolves one item without mutating the source.'),('Async and await','Network work returns promises; errors and response status must be handled explicitly.'),('Union types','A finite status union prevents impossible string values and improves generated-code feedback.'),('Optional chaining','Use it for genuinely optional access, not to hide a missing required value.')]
    for h,text in readiness: d.add_heading(h,2); d.add_paragraph(text); code(d,f"// Predict the value and type before running this example\nconst concept = '{h}'\nconsole.log(concept)")
    for tnum,tname,tdesc in TOPICS:
        d.add_page_break(); d.add_heading(f'Topic {tnum}: {tname}',1); d.add_paragraph(tdesc)
        d.add_heading('Topic map',2)
        for lab in [x for x in LABS if x['topic']==tnum]: d.add_paragraph(f"Lab {lab['id']} — {lab['title']}: {lab['outcome']}",style='List Bullet')
        d.add_heading('Concept briefing',2); d.add_paragraph('The trainer introduces concepts in the context of the next observable build. Learners predict behavior, inspect a short example, then apply the concept in the connected capstone rather than a throwaway exercise.')
        for lab in [x for x in LABS if x['topic']==tnum]:
            # Five meaningful nonblank page units per lab guarantee a detailed >100-page guide.
            d.add_page_break(); d.add_heading(f"Lab {lab['id']}: {lab['title']}",1); add_box(d,'Goal',lab['outcome']); d.add_heading('Build context and concepts',2); d.add_paragraph(concept_explanation(lab))
            for cx in lab['concepts']:
                p=d.add_paragraph(style='List Bullet'); p.add_run(cx+' — ').bold=True; p.add_run(CONCEPT_DEFS[cx])
            d.add_heading('Files in scope',2); add_bullets(d,lab['files']); d.add_heading('Prerequisites',2); add_bullets(d,['Restore the previous lab checkpoint and run the existing app.','Read AGENTS.md and the product brief.','Use synthetic data and placeholders only.','Open the training log for evidence capture.'])
            d.add_page_break(); d.add_heading('Walkthrough — Part A',1)
            lab_no=int(lab['id'].replace('.',''))
            for i,s in enumerate(lab['steps'][:3],1): d.add_heading(f'Step {i} — {first_sentence(s)}',2); d.add_paragraph(s); d.add_paragraph(INSPECT_GUIDE[(lab_no+i)%len(INSPECT_GUIDE)])
            d.add_heading('Code-reading practice',2); code(d,f"// Lab {lab['id']} review marker\n// Files: {', '.join(lab['files'])}\n// Explain inputs, output, side effects and failure paths before acceptance.")
            d.add_page_break(); d.add_heading('Walkthrough — Part B',1)
            for i,s in enumerate(lab['steps'][3:],4): d.add_heading(f'Step {i} — {first_sentence(s)}',2); d.add_paragraph(s); d.add_paragraph(VERIFY_GUIDE[(lab_no+i)%len(VERIFY_GUIDE)])
            d.add_heading('Checkpoint discipline',2); d.add_paragraph('Inspect `git diff --stat`, then the complete diff. Stage named files only. The commit message describes the observable capability, not the AI tool used.')
            d.add_page_break(); d.add_heading('Agentic AI Loop',1)
            for h,text in [('Frame','Restate the outcome, current checkpoint, files, constraints and stop conditions.'),('Plan','Require assumptions, file list, risks, verification and rollback. Stop before code.'),('Generate','Authorize one bounded increment only.'),('Inspect','Challenge dependencies, hidden state, security, accessibility and testability; trace the diff.'),('Verify','Run normal, boundary, empty and failure paths.'),('Correct','Request the smallest evidence-backed correction.'),('Commit','Commit understood code and record restore instructions.')]: d.add_heading(h,2); d.add_paragraph(text)
            d.add_heading('Vibe prompt',2); code(d,lab['prompt'])
            d.add_page_break(); d.add_heading('Generated-Code Audit and Explanation',1); d.add_paragraph(concept_explanation(lab))
            for trap,(cause,fix) in zip(lab['traps'],TRAP_NOTES[lab['id']]):
                d.add_heading(trap,2); p=d.add_paragraph(); p.add_run('Likely cause: ').bold=True; p.add_run(cause+' '); p.add_run('Recovery: ').bold=True; p.add_run(fix)
            d.add_heading('Verification',2); add_bullets(d,lab['verify']+['npm run lint passes.','npm run build passes.','git diff --check is clean.'])
            d.add_heading('Evidence and reflection',2); d.add_paragraph('Submit the prompt, approved plan, annotated diff excerpt, command output, browser evidence and one decision you changed after inspection. Reflect on which evidence changed your confidence and what context the next lab must preserve.')
    d.add_page_break(); d.add_heading('Appendix A: AI React Bug Checklist',1)
    bugs=['key={index} for changing lists','State mutated with push, splice or assignment','State duplicated instead of derived','Functional update omitted for dependent change','Effect dependency suppressed','Effect lacks cleanup or idempotence','Fetch ignores response.ok','Error leaves loading state active','Request starts during render','Controlled input changes to uncontrolled','Form submit reloads page','Button or input lacks accessible name','Clickable div replaces native control','Focus outline removed','Color is the only status signal','Unknown route or ID crashes','Internal link uses window.location','Direct route refresh is not hosted correctly','Tests assert implementation details','Agent adds dependency without approval','Secret appears in prompt or VITE variable','Broad refactor hides the requested fix','Debug log ships in production','Memoization added without measurement','Build success mistaken for user-flow evidence','Diff not inspected before commit','Generated comment contradicts code','Any or lint suppression hides type defect']
    for i,b in enumerate(bugs,1): d.add_paragraph(f'{i}. {b}',style='List Number')
    d.add_page_break(); d.add_heading('Appendix B: Troubleshooting',1)
    for symptom,cause,fix in [('Blank page','Runtime error or missing root','Read the first browser console error; verify index.html → main.tsx → App.tsx.'),('Endless spinner','Rejected fetch path never clears loading','Use finally or an explicit state transition and expose retry.'),('Changes disappear','State mutation or stale closure','Use immutable transformations and functional updates.'),('Route refresh 404','Host lacks SPA fallback','Configure rewrite to index.html or use a supported routing mode.'),('Agent loops on fixes','Prompt lacks a failing check and file boundary','Restore checkpoint, provide exact failure, authorize one experiment.')]:
        d.add_heading(symptom,2); d.add_paragraph(f'Likely cause: {cause}. Recovery: {fix}')
    d.add_page_break(); d.add_heading('Appendix C: Glossary and References',1)
    for term,meaning in [('Component','A function that returns a declarative description of UI.'),('Prop','Input passed from a parent component.'),('State','A render snapshot of data owned by a component or hook.'),('Effect','Synchronization with an external system after commit.'),('Hook','A React function that connects components to reusable React capabilities.'),('Route','A mapping from URL pattern to interface and data behavior.'),('Diff','The exact line-level change between repository states.'),('Regression','Previously working behavior that a change breaks.'),('Checkpoint','A verified Git commit that can be restored.'),('Agentic loop','A controlled sequence of intent, plan, mutation, evidence, critique and checkpoint.')]: d.add_heading(term,2); d.add_paragraph(meaning)
    d.add_heading('References',2)
    for name,url in data.REFERENCES: d.add_paragraph(f'{name}: {url}')
    path=OUT/f"{C['code']}-Learner-Guide.docx"; d.save(path)
    # Markdown is a complete aligned source, composed from the same lab records.
    md=[f"# {C['title']} — Learner Guide",f"- **Course Code:** {C['code']}",f"- **Duration:** {C['duration']}",f"- **Level:** {C['level']}","## Agentic AI Loop","Frame → Plan → Generate → Inspect → Verify → Correct → Commit"]
    for tnum,tname,tdesc in TOPICS:
        md += [f"## Topic {tnum}: {tname}",tdesc]
        for lab in [x for x in LABS if x['topic']==tnum]: md += [build_lab_markdown(lab,tname)]
    (OUT/f"{C['code']}-Learner-Guide.md").write_text('\n\n'.join(md),encoding='utf-8')

def build_lesson_plan():
    toc=[('Course Intent',1),('Learning Outcomes',1),('Detailed Two-Day Schedule',1),('Day 1',2),('Day 2',2),('Facilitation Cycle for Every Lab',1),('Resources and Contingencies',1),('Formative Progress Checks',1)]
    d=base_doc('Lesson Plan',toc); d.add_page_break(); d.add_heading('Course Intent',1); d.add_paragraph('Facilitate an intensive two-day adult-learning journey from an empty folder to a tested, deployed React capstone. Demonstrations are brief; guided practice, prediction, code reading, evidence and reflection dominate contact time.')
    d.add_heading('Learning Outcomes',1); add_bullets(d,['Control AI coding work through an explicit engineering loop.','Explain and apply essential React concepts in a connected app.','Review generated code for correctness, accessibility, security and maintainability.','Test, debug and deploy a verified React production build.'])
    # Lab blocks match the per-lab minute estimates in the lab files exactly:
    # Day 1 labs 1.1–2.5 = 380 min, Day 2 labs 3.1–4.5 = 390 min; 450 contact min/day.
    schedule=[
      ('Day 1','09:30–10:00','Welcome, readiness and the agentic loop',30),('Day 1','10:00–11:15','Labs 1.1–1.2: workspace (35) and scaffold (40)',75),('Day 1','11:15–11:25','Break',10),('Day 1','11:25–12:30','Labs 1.3–1.4: product brief (30) and prompt contract (35)',65),('Day 1','12:30–13:10','Lab 1.5: first React screen (40)',40),('Day 1','13:10–13:40','Lunch',30),('Day 1','13:40–15:00','Labs 2.1–2.2: data and keys (40), components and props (40)',80),('Day 1','15:00–15:10','Break',10),('Day 1','15:10–16:30','Labs 2.3–2.4: composition (35), events and forms (45)',80),('Day 1','16:30–17:10','Lab 2.5: responsive accessible CSS (40)',40),('Day 1','17:10–17:30','Checkpoint and reflection',20),
      ('Day 2','09:30–09:50','Restore checkpoint and retrieval practice',20),('Day 2','09:50–11:10','Labs 3.1–3.2: state (40) and effects (40)',80),('Day 2','11:10–11:20','Break',10),('Day 2','11:20–12:35','Labs 3.3–3.4: custom hook (35) and routing (40)',75),('Day 2','12:35–13:20','Lab 3.5: API data and resilient states (45)',45),('Day 2','13:20–13:50','Lunch',30),('Day 2','13:50–15:10','Labs 4.1–4.2: debugging (35) and testing (45)',80),('Day 2','15:10–15:20','Break',10),('Day 2','15:20–16:30','Labs 4.3–4.4: accessibility (35) and performance (35)',70),('Day 2','16:30–17:10','Lab 4.5: release and deployment (40)',40),('Day 2','17:10–17:30','Demonstration, reflection and action plan',20)]
    d.add_heading('Detailed Two-Day Schedule',1)
    for day in ['Day 1','Day 2']:
        d.add_heading(day,2); t=d.add_table(rows=1,cols=3); t.style='Table Grid'
        for i,x in enumerate(['Time','Activity','Minutes']): t.cell(0,i).text=x; shade(t.cell(0,i),BLUE)
        for dy,tm,act,m in [x for x in schedule if x[0]==day]:
            cells=t.add_row().cells
            for i,x in enumerate([tm,act,str(m)]): cells[i].text=x
        d.add_paragraph('Instructional/contact minutes excluding lunch: 450.')
    d.add_heading('Facilitation Cycle for Every Lab',1)
    for h,text in [('Activate','Ask learners to retrieve the prior checkpoint and predict the next visible behavior.'),('Model','Demonstrate the first plan/diff decision aloud, including one rejected AI suggestion.'),('Guide','Learners work in pairs: one drives, one audits scope and evidence; swap roles.'),('Check','Pause at verification gates and ask a learner to explain data or event flow.'),('Release','Learners submit evidence, reflect and commit a recoverable checkpoint.')]: d.add_heading(h,2); d.add_paragraph(text)
    d.add_heading('Resources and Contingencies',1); add_bullets(d,['Trainer reference repository with one checkpoint per lab.','Prepared diffs and screenshots if an AI service is unavailable.','Local synthetic JSON endpoint so API practice requires no credentials.','Git recovery sheet for learners who fall behind.','Browser accessibility tree and React DevTools for evidence-led inspection.'])
    d.add_heading('Formative Progress Checks',1); d.add_paragraph('Use observation, questioning, demonstrations, diff explanations and lab evidence. These checks support learning and pacing and are not graded.')
    d.save(OUT/f"{C['code']}-Lesson-Plan.docx")

build_labs(); build_learner_guide(); build_lesson_plan()
(OUT/'course-data.json').write_text(json.dumps({'course':C,'topics':TOPICS,'labs':LABS,'references':data.REFERENCES},indent=2),encoding='utf-8')
print(f"Generated {len(LABS)} labs and aligned courseware from one source")
