#!/usr/bin/env python3
"""Apply a richer, consistent visual system to the C1143 facilitator deck."""
from pathlib import Path
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "courseware" / "C1143-Facilitator-Deck.pptx"
OUT = ROOT / "courseware" / "C1143-Facilitator-Deck-Visual-Enhanced.pptx"

INK = "142033"
MUTED = "5B6372"
PALE = "F5F8FC"
WHITE = "FFFFFF"
TOPIC = {0: "0B6E99", 1: "0B8F87", 2: "7C3AED", 3: "1F6FEB", 4: "C65D21"}
SOFT = {0: "EAF4F8", 1: "E8F7F4", 2: "F1EBFD", 3: "EAF2FE", 4: "FCEFE7"}


def rgb(value):
    return RGBColor.from_string(value)


def send_to_back(shape):
    tree = shape._element.getparent()
    tree.remove(shape._element)
    tree.insert(2, shape._element)


def add_rect(slide, x, y, w, h, color, radius=False, line=None, back=True):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(color)
    sh.line.color.rgb = rgb(line or color)
    if back: send_to_back(sh)
    return sh


def add_circle(slide, x, y, d, color, back=True):
    sh = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(color); sh.line.fill.background()
    if back: send_to_back(sh)
    return sh


def add_text(slide, text, x, y, w, h, size, color, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; r.font.name = "Arial"; r.font.size = Pt(size)
    r.font.bold = bold; r.font.color.rgb = rgb(color)
    return box


def slide_text(slide):
    return "\n".join(s.text for s in slide.shapes if hasattr(s, "text") and s.text)


def normalize_loop_labels(slide):
    """Align standalone process labels with the non-WSQ seven-stage loop."""
    mapping = {
        "Specify": "Frame", "Inspect": "Generate", "Implement": "Inspect",
        "Test": "Verify", "Critique": "Correct", "Refine": "Commit",
        "Checkpoint": "Reflect",
    }
    for sh in slide.shapes:
        if not hasattr(sh, "text_frame"): continue
        for p in sh.text_frame.paragraphs:
            if p.text.strip() in mapping and len(p.runs) == 1:
                p.runs[0].text = mapping[p.text.strip()]


def topic_for(text):
    m = re.search(r"(?:TOPIC|Topic)\s*([1-4])", text)
    if m: return int(m.group(1))
    m = re.search(r"LAB\s*([1-4])\.", text)
    return int(m.group(1)) if m else 0


def lab_stage(text):
    titles = [
        "Prepare", "Concepts", "Sequence", "Prompt", "Audit", "Verify"
    ]
    if "Concepts and observable outcome" in text: return 1, titles[1]
    if "Executable lab sequence" in text: return 2, titles[2]
    if "Vibe prompt" in text: return 3, titles[3]
    if "Read what the AI wrote" in text: return 4, titles[4]
    if "Verification and checkpoint" in text: return 5, titles[5]
    if re.search(r"LAB\s*[1-4]\.\d", text): return 0, titles[0]
    return None, None


def add_dot_grid(slide, accent):
    for row in range(3):
        for col in range(5):
            add_circle(slide, 12.48 + col * .12, 6.86 + row * .12, .035, accent)


def restyle_existing_cards(slide, accent, soft):
    """Give existing filled cards a little more hierarchy without moving content."""
    cards = []
    for sh in slide.shapes:
        if not hasattr(sh, "fill") or not sh.fill.type or not hasattr(sh, "text_frame"):
            continue
        if sh.width < Inches(1.35) or sh.height < Inches(.55):
            continue
        if sh.top < Inches(.75) or sh.top > Inches(6.6):
            continue
        cards.append(sh)
    for i, sh in enumerate(cards):
        # Preserve dark code/prompt panels.
        try:
            current = str(sh.fill.fore_color.rgb)
        except Exception:
            current = ""
        if current in {"142033", "161B26", "17212B"}:
            continue
        if i % 3 == 0:
            sh.fill.solid(); sh.fill.fore_color.rgb = rgb(soft)
            sh.line.color.rgb = rgb(accent)
            sh.line.width = Pt(1.15)


def decorate_content(slide, idx, text, topic):
    accent, soft = TOPIC[topic], SOFT[topic]
    # Pale geometry makes the canvas feel intentional while staying readable.
    add_circle(slide, 11.58, -.55, 2.15, soft)
    add_circle(slide, 12.42, .12, .38, accent)
    add_rect(slide, 0, 7.32, 13.333, .18, soft)
    add_rect(slide, 0, 0, .08, 7.5, accent)
    add_dot_grid(slide, accent)

    stage, label = lab_stage(text)
    if stage is not None:
        # Persistent six-stage lab navigator.
        x0, gap, segw = 8.92, .08, .60
        for s in range(6):
            c = accent if s == stage else "DDE5EE"
            add_rect(slide, x0 + s * (segw + gap), .24, segw, .055, c, radius=True)
        add_text(slide, label.upper(), 11.85, .33, 1.0, .22, 7.5, accent, True, PP_ALIGN.RIGHT)

        m = re.search(r"LAB\s*([1-4]\.\d)", text)
        if m:
            add_text(slide, m.group(1), 11.65, 5.75, 1.25, .72, 35, soft, True, PP_ALIGN.RIGHT)
    else:
        add_text(slide, f"{idx:03}", 11.95, 6.36, .75, .26, 8, accent, True, PP_ALIGN.RIGHT)
    restyle_existing_cards(slide, accent, soft)


def decorate_section(slide, text, topic):
    accent, soft = TOPIC[topic], SOFT[topic]
    add_rect(slide, 0, 0, .22, 7.5, accent)
    add_rect(slide, 8.65, 0, 4.683, 7.5, soft)
    # Abstract component tree / agent flow on the right.
    nodes = [(9.15,5.52),(10.78,5.52),(12.05,5.52)]
    for x,y in nodes:
        add_rect(slide, x, y, 1.15, .72, WHITE, radius=True, line=accent, back=False)
    for x,y,w,h in [(10.30,5.84,.48,.08),(11.93,5.84,.12,.08)]:
        add_rect(slide,x,y,w,h,accent,back=False)
    add_text(slide, "</>  CODE", 9.32, 5.70, .82, .24, 9, accent, True, PP_ALIGN.CENTER)
    add_text(slide, "AI  PLAN", 10.94, 5.70, .82, .24, 9, accent, True, PP_ALIGN.CENTER)
    add_text(slide, "SHIP", 12.22, 5.70, .70, .24, 9, accent, True, PP_ALIGN.CENTER)
    add_dot_grid(slide, accent)


def decorate_cover(slide, closing=False):
    accent, soft = TOPIC[3], SOFT[3]
    add_rect(slide, 0, 0, .24, 7.5, accent)
    add_circle(slide, 9.6, -.85, 4.6, soft)
    add_circle(slide, 10.65, .4, 2.3, WHITE, back=False)
    add_circle(slide, 11.20, .95, 1.2, accent, back=False)
    add_text(slide, "{ }", 10.87, 1.28, 1.85, .72, 34, WHITE, True, PP_ALIGN.CENTER)
    for i, word in enumerate(["SPECIFY", "PLAN", "BUILD", "VERIFY"]):
        add_rect(slide, 9.18, 3.30 + i*.56, 2.75, .38, WHITE, radius=True, line=accent, back=False)
        add_text(slide, f"0{i+1}  {word}", 9.42, 3.37+i*.56, 2.18, .22, 9, accent, True)
    add_dot_grid(slide, accent)


def main():
    prs = Presentation(SRC)
    for idx, slide in enumerate(prs.slides, 1):
        normalize_loop_labels(slide)
        text = slide_text(slide)
        topic = topic_for(text)
        if idx in (1, len(prs.slides)):
            decorate_cover(slide, idx == len(prs.slides))
        elif re.search(r"Topic [1-4]:", text) and "TERTIARY INFOTECH ACADEMY" in text:
            decorate_section(slide, text, topic)
        else:
            decorate_content(slide, idx, text, topic)
    prs.save(OUT)
    print(f"Saved {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
