#!/usr/bin/env python3
"""Generate the C1143 Lesson Plan (LP) DOCX in the Tertiary house format (non-WSQ: no assessment).

Cover page + Document Version Control Record + auto TOC + Arial 11pt body +
colour-coded 2-day schedule tables (9:00am-6:00pm, 8 training hours/day, 1h
lunch, tea breaks within training time; no assessment — the afternoon closes with the mini-capstone).
Topics/labs come from course_data + the domain data files so the LP stays
aligned with the deck, the Learner Guide and the labs.
"""
import os, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE)
import course_data as C
from data_domain1 import DOMAIN1; from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3; from data_domain4 import DOMAIN4
from data_domain5 import DOMAIN5; from data_domain6 import DOMAIN6
ACT=DOMAIN1+DOMAIN2+DOMAIN3+DOMAIN4+DOMAIN5+DOMAIN6
import prodoc
REPO=os.path.dirname(os.path.dirname(HERE)); ASSETS=os.path.join(REPO,"courseware","assets")

BRAND=RGBColor(0x1F,0x6F,0xEB); GREY=RGBColor(0x55,0x5B,0x66)
HEADER_FILL="1F6FEB"; TOPIC_FILL="E8F0FE"; BREAK_FILL="FFF4E5"; LUNCH_FILL="FDE9D9"; ASSESS_FILL="E8F7EE"

def lab_titles(nums):
    return "; ".join(f"Lab {a['num']}: {a['title']}" for a in ACT if a['num'] in nums)

# ------------------------------------------------ schedule (single source of truth for timing)
# (start, end, minutes, kind, activity_text)  kind: admin/topic/lab/break/lunch/assess/recap
# Every day totals exactly 480 training minutes (the 1-hour lunch is excluded).
SCHEDULE = {
 1: (C.DAY_THEMES[1], [
    ("9:00","9:30",30,"admin","Welcome, course introduction and ground rules"),
    ("9:30","10:30",60,"topic","Topic 1 — Build a React Web App Using Vibe Coding: the Prompt → Generate → Read → Understand → Correct loop, JSX, components and props (concepts + demo)"),
    ("10:30","10:45",15,"break","Tea break"),
    ("10:45","12:30",105,"lab","Hands-on: "+lab_titles(["1.1","1.2","1.3"])),
    ("12:30","13:00",30,"topic","Topic 2 — Deploy Your React Web App to the Cloud: Git as a safety net, the production build, CI/CD and environment variables (concepts)"),
    ("13:00","14:00",60,"lunch","Lunch break"),
    ("14:00","15:15",75,"lab","Hands-on: "+lab_titles(["2.1","2.2","2.3"])),
    ("15:15","15:30",15,"break","Tea break"),
    ("15:30","16:15",45,"topic","Topic 3 — Improving Your App by Learning Core React Concepts: the Virtual DOM, composition, lists and keys, events and controlled forms (concepts + demo)"),
    ("16:15","17:45",90,"lab","Hands-on: "+lab_titles(["3.1","3.2","3.3","3.4"])),
    ("17:45","18:00",15,"recap","Day 1 recap and Q&A"),
 ]),
 2: (C.DAY_THEMES[2], [
    ("9:00","9:15",15,"recap","Day 1 recap and knowledge check-in"),
    ("9:15","9:45",30,"topic","Topic 4 — React Hooks (The Vibe Way): the Rules of Hooks, state as a snapshot, immutable updates and effect cleanup (concepts + demo)"),
    ("9:45","11:00",75,"lab","Hands-on: "+lab_titles(["4.1","4.2","4.3","4.4","4.5"])),
    ("11:00","11:10",10,"break","Tea break"),
    ("11:10","11:40",30,"topic","Topic 5 — Giving Your App a Backend API and a Database: the three-tier model, serverless API routes over Neon, parameterised SQL, the three states of a remote read, and bcrypt + JWT authentication (concepts)"),
    ("11:40","13:00",80,"lab","Hands-on: "+lab_titles(["5.1","5.2","5.3","5.4"])),
    ("13:00","14:00",60,"lunch","Lunch break"),
    ("14:00","14:20",20,"topic","Topic 6 — React Router for Real App Navigation: single page applications, layout routes, dynamic parameters and protected routes (concepts)"),
    ("14:20","15:10",50,"lab","Hands-on: "+lab_titles(["6.1","6.2","6.3"])),
    ("15:10","15:20",10,"break","Tea break"),
    ("15:20","16:20",60,"lab","Hands-on: "+lab_titles(["6.4","6.5"])),
    ("16:20","17:40",80,"lab","Hands-on: "+lab_titles(["6.6"])+" — the mini-capstone: add a whole course-reviews feature yourself, mostly by prompting"),
    ("17:40","18:00",20,"recap","Course wrap-up, next steps and Q&A"),
 ]),
}

# ------------------------------------------------ build document
doc=Document()
normal=doc.styles["Normal"]; normal.font.name="Arial"; normal.font.size=Pt(11)
prodoc.style_headings(doc)

prodoc.add_cover_page(doc,"LESSON PLAN",C.TITLE,C.VERSION.lstrip("v"),
                      org_logo=os.path.join(ASSETS,"tertiary-infotech-logo.png"),
                      course_logo=None, course_code=C.COURSE_CODE)
prodoc.add_version_control(doc,[
    ("2.0","12 July 2026","C1143 20-lab agentic-loop edition.",C.TRAINER),
    ("2.1","21 July 2026","Title, contents and schedule corrections.",C.TRAINER),
    (C.VERSION.lstrip("v"),C.VERSION_DATE,
     "Rebuilt on the flagship 2-day Full Stack React with Vibe Coding courseware: 25 hands-on labs across six topics "
     "building the Cook & Bake Academy app end to end. Non-assessed delivery — the final afternoon is reallocated to "
     "the deployment lab and the mini-capstone.",C.TRAINER)])
prodoc.add_toc(doc)

def H(text,level=1):
    return doc.add_heading(text,level=level)

H("Course Information",1)
info=[("Course Title",C.TITLE),("Course Code",C.COURSE_CODE),
      ("Training Provider",C.ORG+"  ("+C.UEN.replace('UEN: ','UEN ')+")"),
      ("Duration",f"{C.DAYS} days · 8 training hours per day (16 hours)"),
      ("Daily Timing","9:00 am – 6:00 pm (1-hour lunch; tea breaks within training time)"),
      ("Mode","Instructor-led, hands-on labs in VS Code with an AI coding agent"),
      ("Project",f"{C.PROJECT} — a full-stack React course catalogue built across all six topics"),
      ("Technology","Vite, React 19, React Router 7, Vercel serverless functions, Neon Postgres, Vercel"),
      ("Trainer",C.TRAINER)]
t=doc.add_table(rows=0,cols=2); t.style="Table Grid"
for k,v in info:
    c=t.add_row().cells; c[0].text=""; r=c[0].paragraphs[0].add_run(k); r.bold=True; r.font.size=Pt(10)
    prodoc._shade_cell(c[0],TOPIC_FILL)
    c[1].text=""; c[1].paragraphs[0].add_run(v).font.size=Pt(10)

H("Learning Outcomes",1)
doc.add_paragraph("On completion of this course, learners will be able to:")
for lo in C.LEARNING_OUTCOMES:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(lo).font.size=Pt(10.5)

H("Learning Reinforcement",1)
for a in ["This is a non-assessed short course: there is no written or practical assessment.",
          "Every lab ends with an observable 'Test it' check the learner verifies in the browser before moving on.",
          "The trainer checks understanding through questioning, code walk-throughs and review of each learner's running app.",
          "The Day 2 afternoon is a capstone block: learners deploy the full-stack app and then add a complete feature themselves in the mini-capstone (Lab 6.6).",
          "Each day closes with a recap and Q&A; every lab folder is a restorable checkpoint so no learner is left behind."]:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(a).font.size=Pt(10.5)

def set_cell(cell,text,bold=False,size=9.5,color=None,fill=None,align=None):
    cell.text=""; p=cell.paragraphs[0]
    if align: p.alignment=align
    r=p.add_run(text); r.bold=bold; r.font.size=Pt(size); r.font.name="Arial"
    if color: r.font.color.rgb=color
    if fill: prodoc._shade_cell(cell,fill)

KIND_FILL={"topic":TOPIC_FILL,"break":BREAK_FILL,"lunch":LUNCH_FILL,"assess":ASSESS_FILL,
           "admin":"F3F5F8","recap":"F3F5F8","lab":None}

H("Course Schedule",1)
for day,(theme,rows) in SCHEDULE.items():
    H(f"Day {day} — {theme}",2)
    tbl=doc.add_table(rows=0,cols=3); tbl.style="Table Grid"; tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
    hdr=tbl.add_row().cells
    for i,htext in enumerate(["Time","Duration","Topic / Activity"]):
        set_cell(hdr[i],htext,bold=True,size=10,color=RGBColor(0xFF,0xFF,0xFF),fill=HEADER_FILL)
    training=0
    for start,end,mins,kind,text in rows:
        cells=tbl.add_row().cells; fill=KIND_FILL.get(kind)
        set_cell(cells[0],f"{start}–{end}",bold=(kind in ("topic","assess")),size=9.5,fill=fill)
        set_cell(cells[1],f"{mins} min",size=9.5,fill=fill)
        set_cell(cells[2],text,bold=(kind in ("topic","assess")),size=9.5,fill=fill)
        if kind!="lunch": training+=mins
    for row in tbl.rows:
        row.cells[0].width=Inches(1.15); row.cells[1].width=Inches(0.9); row.cells[2].width=Inches(4.75)
    p=doc.add_paragraph(); r=p.add_run(f"Total training time: {training} minutes ({training//60} hours).")
    r.italic=True; r.font.size=Pt(9.5); r.font.color.rgb=GREY
    assert training==480, f"Day {day} training minutes = {training}, expected 480"

H("Lab Reference (aligned to the course topics)",1)
tt=doc.add_table(rows=0,cols=3); tt.style="Table Grid"
hdr=tt.add_row().cells
for i,htext in enumerate(["Topic","Scope","Labs"]):
    set_cell(hdr[i],htext,bold=True,size=10,color=RGBColor(0xFF,0xFF,0xFF),fill=HEADER_FILL)
for tp in C.TOPICS:
    acts=[a for a in ACT if a["topic"]==tp["num"]]
    cells=tt.add_row().cells
    set_cell(cells[0],f"Topic {tp['code']}: {tp['title']}",bold=True,size=9.5,fill=TOPIC_FILL)
    set_cell(cells[1],tp["weighting"],size=9.5,fill=TOPIC_FILL)
    set_cell(cells[2],", ".join(f"Lab {a['num']}" for a in acts),size=9.5)

prodoc.add_page_numbers(doc)
prodoc.enable_update_fields(doc)
OUT=os.path.join(REPO,"courseware",f"LP-{C.SHORT_TITLE}.docx")
doc.save(OUT)
print("Saved",OUT)
